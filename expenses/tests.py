from datetime import date
from decimal import Decimal
from io import BytesIO
from openpyxl import load_workbook
from django.contrib.auth.models import User
from django.test import TestCase
from django.utils import timezone
from django.urls import reverse
from .models import Budget, Transaction

class TrackEZTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="alice", password="StrongPass123")
        self.other = User.objects.create_user(username="bob", password="StrongPass123")
        self.client.login(username="alice", password="StrongPass123")

    def add(self, title, amount, kind, category=Transaction.OTHER, day=date(2026, 1, 1), user=None):
        return Transaction.objects.create(
            user=user or self.user, title=title, amount=amount,
            transaction_type=kind, category=category, date=day
        )

    def test_dashboard_totals(self):
        self.add("Salary", "50000.00", Transaction.INCOME)
        self.add("Rent", "15000.00", Transaction.EXPENSE)
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.context["income"], Decimal("50000.00"))
        self.assertEqual(response.context["expense"], Decimal("15000.00"))
        self.assertEqual(response.context["balance"], Decimal("35000.00"))

    def test_add_income_and_expense(self):
        self.client.post(reverse("add_income"), {
            "title": "Freelance", "amount": "2500", "category": "Other", "date": "2026-02-10"
        })
        self.client.post(reverse("add_expense"), {
            "title": "Travel", "amount": "800", "category": "Travel", "date": "2026-02-10"
        })
        self.assertEqual(Transaction.objects.count(), 2)
        self.assertEqual(Transaction.objects.get(title="Freelance").transaction_type, Transaction.INCOME)
        self.assertEqual(Transaction.objects.get(title="Travel").transaction_type, Transaction.EXPENSE)
        self.assertEqual(Transaction.objects.get(title="Travel").category, Transaction.TRAVEL)

    def test_user_isolation(self):
        self.add("Private", "9000", Transaction.INCOME, user=self.other)
        self.add("Mine", "1000", Transaction.INCOME)
        response = self.client.get(reverse("history"))
        self.assertContains(response, "Mine")
        self.assertNotContains(response, "Private")

    def test_delete_cannot_delete_other_user_data(self):
        mine = self.add("Mine", "100", Transaction.EXPENSE)
        theirs = self.add("Theirs", "200", Transaction.EXPENSE, user=self.other)
        self.client.post(reverse("delete_transaction", args=[theirs.pk]))
        self.assertTrue(Transaction.objects.filter(pk=theirs.pk).exists())
        self.client.post(reverse("delete_transaction", args=[mine.pk]))
        self.assertFalse(Transaction.objects.filter(pk=mine.pk).exists())

    def test_quick_transaction_does_not_save(self):
        response = self.client.post(reverse("quick"), {"income": "1000", "expense": ["100", "50", "25"]})
        self.assertContains(response, "825")
        self.assertEqual(Transaction.objects.count(), 0)

    def test_history_date_filter(self):
        self.add("January", "100", Transaction.EXPENSE, day=date(2026, 1, 10))
        self.add("February", "200", Transaction.EXPENSE, day=date(2026, 2, 10))
        response = self.client.get(reverse("history"), {"start": "2026-02-01", "end": "2026-02-28"})
        self.assertContains(response, "February")
        self.assertNotContains(response, "January")

    def test_excel_export(self):
        self.add("February", "200", Transaction.EXPENSE, category=Transaction.TRAVEL, day=date(2026, 2, 10))
        response = self.client.get(reverse("export"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Title", "Amount", "Type", "Category", "Date"))
        self.assertEqual(rows[1][0], "February")
        self.assertEqual(rows[1][3], "Travel")

    def test_budget_creation_and_dashboard_alert(self):
        current_month = timezone.localdate().replace(day=1)
        Budget.objects.create(user=self.user, month=current_month, category=Transaction.FOOD, amount="1000")
        self.add("Food", "900", Transaction.EXPENSE, category=Transaction.FOOD, day=timezone.localdate())
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(any("used 90%" in x for x in response.context["insights"]))

    def test_budget_page_and_monthly_report(self):
        Budget.objects.create(user=self.user, month=date(2026, 1, 1), category=Transaction.TRAVEL, amount="2000")
        self.add("Trip", "1500", Transaction.EXPENSE, category=Transaction.TRAVEL, day=date(2026, 1, 15))
        self.add("Salary", "5000", Transaction.INCOME, day=date(2026, 1, 1))
        self.assertEqual(self.client.get(reverse("budgets")).status_code, 200)
        response = self.client.get(reverse("monthly_report"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["rows"][0]["expense"], Decimal("1500"))
        self.assertEqual(response.context["rows"][0]["income"], Decimal("5000"))

    def test_login_required(self):
        self.client.logout()
        for name in ["dashboard", "add_income", "add_expense", "history", "export", "quick", "budgets", "monthly_report"]:
            response = self.client.get(reverse(name))
            self.assertEqual(response.status_code, 302)
            self.assertIn("/login/", response["Location"])
