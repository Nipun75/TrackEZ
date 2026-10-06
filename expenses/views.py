import json
from datetime import date
from decimal import Decimal, InvalidOperation
from django.contrib import messages
from django.contrib.auth import login
from django.contrib.auth.decorators import login_required
from django.db.models import Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from .forms import BudgetForm, RegisterForm, TransactionForm
from .models import Budget, Transaction

def register(request):
    if request.user.is_authenticated:
        return redirect("dashboard")
    form = RegisterForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        user = form.save()
        login(request, user)
        messages.success(request, "Account created successfully.")
        return redirect("dashboard")
    return render(request, "register.html", {"form": form})

def _month_start(value):
    return value.replace(day=1)

@login_required
def dashboard(request):
    qs = Transaction.objects.filter(user=request.user)
    today = timezone.localdate()
    month_start = today.replace(day=1)
    month_qs = qs.filter(date__gte=month_start, date__lte=today)
    income = qs.filter(transaction_type=Transaction.INCOME).aggregate(v=Sum("amount"))["v"] or Decimal("0")
    expense = qs.filter(transaction_type=Transaction.EXPENSE).aggregate(v=Sum("amount"))["v"] or Decimal("0")
    month_income = month_qs.filter(transaction_type=Transaction.INCOME).aggregate(v=Sum("amount"))["v"] or Decimal("0")
    month_expense = month_qs.filter(transaction_type=Transaction.EXPENSE).aggregate(v=Sum("amount"))["v"] or Decimal("0")

    dates = sorted({x.date.isoformat() for x in qs})
    income_by_date = {d: sum((x.amount for x in qs if x.date.isoformat() == d and x.transaction_type == Transaction.INCOME), Decimal("0")) for d in dates}
    expense_by_date = {d: sum((x.amount for x in qs if x.date.isoformat() == d and x.transaction_type == Transaction.EXPENSE), Decimal("0")) for d in dates}

    top_expenses = qs.filter(transaction_type=Transaction.EXPENSE).order_by("-amount")[:5]
    top_income = qs.filter(transaction_type=Transaction.INCOME).order_by("-amount")[:5]

    category_rows = []
    for category, label in Transaction.CATEGORY_CHOICES:
        spent = month_qs.filter(transaction_type=Transaction.EXPENSE, category=category).aggregate(v=Sum("amount"))["v"] or Decimal("0")
        budget = Budget.objects.filter(user=request.user, month=month_start, category=category).first()
        if spent or budget:
            budget_amount = budget.amount if budget else Decimal("0")
            usage = (spent / budget_amount * 100) if budget_amount else Decimal("0")
            category_rows.append({"category": label, "spent": spent, "budget": budget_amount, "usage": min(usage, Decimal("999"))})

    insights = []
    if month_income:
        savings_rate = (month_income - month_expense) / month_income * 100
        insights.append(f"Your current-month savings rate is {savings_rate:.1f}%.")
    if category_rows:
        highest = max(category_rows, key=lambda x: x["spent"])
        if highest["spent"] > 0:
            insights.append(f"{highest['category']} is your highest current-month expense category at ₹{highest['spent']:.2f}.")
        for row in category_rows:
            if row["budget"] and row["spent"] > row["budget"]:
                insights.append(f"Budget alert: {row['category']} is over budget by ₹{row['spent'] - row['budget']:.2f}.")
            elif row["budget"] and row["usage"] >= 80:
                insights.append(f"Budget alert: {row['category']} has used {row['usage']:.0f}% of its budget.")

    return render(request, "dashboard.html", {
        "income": income, "expense": expense, "balance": income - expense, "recent": qs[:5],
        "month_income": month_income, "month_expense": month_expense,
        "labels_json": json.dumps(dates),
        "income_data_json": json.dumps([float(income_by_date[d]) for d in dates]),
        "expense_data_json": json.dumps([float(expense_by_date[d]) for d in dates]),
        "top_expense_labels_json": json.dumps([x.title for x in top_expenses]),
        "top_expense_data_json": json.dumps([float(x.amount) for x in top_expenses]),
        "top_income_labels_json": json.dumps([x.title for x in top_income]),
        "top_income_data_json": json.dumps([float(x.amount) for x in top_income]),
        "category_rows": category_rows, "insights": insights,
        "month_label": today.strftime("%B %Y"),
    })

@login_required
def add_income(request):
    return transaction_form(request, Transaction.INCOME)

@login_required
def add_expense(request):
    return transaction_form(request, Transaction.EXPENSE)

def transaction_form(request, kind):
    form = TransactionForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        obj = form.save(commit=False)
        obj.user = request.user
        obj.transaction_type = kind
        obj.save()
        messages.success(request, f"{kind.title()} added successfully.")
        return redirect("dashboard")
    return render(request, "transaction_form.html", {"form": form, "kind": kind.title()})

@login_required
def history(request):
    qs = Transaction.objects.filter(user=request.user)
    start_date, end_date = request.GET.get("start", ""), request.GET.get("end", "")
    category = request.GET.get("category", "")
    transaction_type = request.GET.get("type", "")
    if start_date:
        try: qs = qs.filter(date__gte=date.fromisoformat(start_date))
        except ValueError: messages.error(request, "Invalid start date."); start_date = ""
    if end_date:
        try: qs = qs.filter(date__lte=date.fromisoformat(end_date))
        except ValueError: messages.error(request, "Invalid end date."); end_date = ""
    if category in dict(Transaction.CATEGORY_CHOICES):
        qs = qs.filter(category=category)
    if transaction_type in (Transaction.INCOME, Transaction.EXPENSE):
        qs = qs.filter(transaction_type=transaction_type)
    return render(request, "history.html", {
        "transactions": qs, "start_date": start_date, "end_date": end_date,
        "categories": Transaction.CATEGORY_CHOICES, "selected_category": category,
        "selected_type": transaction_type,
    })

@login_required
def delete_transaction(request, pk):
    obj = get_object_or_404(Transaction, pk=pk, user=request.user)
    if request.method == "POST":
        obj.delete()
        messages.success(request, "Transaction deleted.")
    return redirect("history")

@login_required
def budget_list(request):
    if request.method == "POST":
        form = BudgetForm(request.POST)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.user = request.user
            obj.month = _month_start(obj.month)
            Budget.objects.update_or_create(user=request.user, month=obj.month, category=obj.category, defaults={"amount": obj.amount})
            messages.success(request, "Monthly budget saved.")
            return redirect("budgets")
    else:
        form = BudgetForm(initial={"month": timezone.localdate().replace(day=1)})
    today = timezone.localdate()
    month_start = today.replace(day=1)
    budgets = Budget.objects.filter(user=request.user, month=month_start)
    rows = []
    for budget in budgets:
        spent = Transaction.objects.filter(user=request.user, date__gte=month_start, date__lte=today, transaction_type=Transaction.EXPENSE, category=budget.category).aggregate(v=Sum("amount"))["v"] or Decimal("0")
        usage = spent / budget.amount * 100 if budget.amount else Decimal("0")
        rows.append({"budget": budget, "spent": spent, "remaining": budget.amount - spent, "usage": usage})
    return render(request, "budgets.html", {"form": form, "rows": rows, "month_label": today.strftime("%B %Y")})

@login_required
def delete_budget(request, pk):
    obj = get_object_or_404(Budget, pk=pk, user=request.user)
    if request.method == "POST":
        obj.delete()
        messages.success(request, "Budget deleted.")
    return redirect("budgets")

@login_required
def monthly_report(request):
    qs = Transaction.objects.filter(user=request.user)
    rows = []
    months = sorted({_month_start(t.date) for t in qs}, reverse=True)
    for month in months:
        income = qs.filter(date__year=month.year, date__month=month.month, transaction_type=Transaction.INCOME).aggregate(v=Sum("amount"))["v"] or Decimal("0")
        expense = qs.filter(date__year=month.year, date__month=month.month, transaction_type=Transaction.EXPENSE).aggregate(v=Sum("amount"))["v"] or Decimal("0")
        rows.append({"month": month, "income": income, "expense": expense, "balance": income - expense})
    return render(request, "monthly_report.html", {"rows": rows})

@login_required
def export_transactions(request):
    qs = Transaction.objects.filter(user=request.user)
    for key, lookup in (("start", "date__gte"), ("end", "date__lte")):
        value = request.GET.get(key)
        if value:
            try: qs = qs.filter(**{lookup: date.fromisoformat(value)})
            except ValueError: return HttpResponse(f"Invalid {key} date.", status=400)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transactions"
    headers = ["Title", "Amount", "Type", "Category", "Date"]
    sheet.append(headers)
    for cell in sheet[1]: cell.font = Font(bold=True)
    for t in qs: sheet.append([t.title, float(t.amount), t.get_transaction_type_display(), t.get_category_display(), t.date])
    for i, width in enumerate([28, 16, 16, 20, 16], 1): sheet.column_dimensions[get_column_letter(i)].width = width
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="trackez_transactions.xlsx"'
    workbook.save(response)
    return response

@login_required
def quick_transaction(request):
    result = None
    if request.method == "POST":
        try:
            income = Decimal(request.POST.get("income", "0") or "0")
            expenses = sum((Decimal(x or "0") for x in request.POST.getlist("expense")), Decimal("0"))
            result = income - expenses
        except InvalidOperation:
            messages.error(request, "Please enter valid numbers.")
    return render(request, "quick.html", {"result": result})
